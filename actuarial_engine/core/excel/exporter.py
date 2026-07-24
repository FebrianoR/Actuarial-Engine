"""
Excel Exporter – Generate hasil kalkulasi PSAK 117 ke file Excel

Format output: 
  Sheet 1: Ringkasan (Summary)
  Sheet 2: Detail per polis
  Sheet 3: Sistem Pembukuan GMM (seperti format template)
"""
from __future__ import annotations

import io
from datetime import datetime, date
from typing import Any

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter


# Warna tema PSAK 117
HEADER_BLUE = "1B3A6B"
HEADER_LIGHT = "2E6DB4"
ROW_ALT = "EBF3FB"
ACCENT_GREEN = "2E7D32"
ACCENT_RED = "C62828"
WHITE = "FFFFFF"


def _fmt_currency(ws, cell_ref: str) -> None:
    """Format cell sebagai IDR."""
    ws[cell_ref].number_format = '#,##0'


def _header_style(cell, bold: bool = True, color: str = HEADER_BLUE) -> None:
    cell.font = Font(bold=bold, color=WHITE, size=10)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _border_all(cell) -> None:
    thin = Side(style="thin", color="CCCCCC")
    cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def export_results(batch_results: list[dict[str, Any]], valuation_date: str | None = None) -> bytes:
    """
    Generate Excel file dari hasil kalkulasi batch.
    
    Args:
        batch_results: List dict hasil kalkulasi (dari BatchService)
        valuation_date: Tanggal valuasi
    
    Returns:
        bytes: File Excel dalam format .xlsx
    """
    wb = openpyxl.Workbook()

    # ── Sheet 1: Summary ──────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Ringkasan"

    val_date = valuation_date or datetime.now().strftime("%d-%m-%Y")
    total_policies = len(batch_results)
    total_lrc = sum(r.get("lrc_net", r.get("icl", 0)) for r in batch_results)
    total_lic = sum(r.get("lic", 0) for r in batch_results)
    total_icl = sum(r.get("insurance_contract_liability", r.get("total_liability", 0)) for r in batch_results)

    # Title
    ws_sum.merge_cells("A1:H1")
    ws_sum["A1"] = f"LAPORAN PSAK 117 – HASIL KALKULASI AKTUARIA"
    ws_sum["A1"].font = Font(bold=True, size=14, color=HEADER_BLUE)
    ws_sum["A1"].alignment = Alignment(horizontal="center")

    ws_sum.merge_cells("A2:H2")
    ws_sum["A2"] = f"Tanggal Valuasi: {val_date}  |  Total Polis: {total_policies:,}"
    ws_sum["A2"].alignment = Alignment(horizontal="center")

    ws_sum.row_dimensions[3].height = 5

    # Summary cards
    headers_sum = ["Metrik", "Nilai (IDR)"]
    summary_data = [
        ("Total Polis", total_policies),
        ("Total LRC Neto", total_lrc),
        ("Total LIC (Klaim)", total_lic),
        ("Total ICL (Liabilitas)", total_icl),
        ("PAA Policies", sum(1 for r in batch_results if r.get("method") == "PAA")),
        ("GMM Policies", sum(1 for r in batch_results if r.get("method") == "GMM")),
        ("VFA Policies", sum(1 for r in batch_results if r.get("method") == "VFA")),
    ]

    for col, h in enumerate(headers_sum, 2):
        cell = ws_sum.cell(row=4, column=col, value=h)
        _header_style(cell)

    for row_i, (label, val) in enumerate(summary_data, 5):
        ws_sum.cell(row=row_i, column=2, value=label)
        cell_val = ws_sum.cell(row=row_i, column=3, value=val if isinstance(val, int) else round(val, 2))
        if isinstance(val, float):
            cell_val.number_format = '#,##0.00'
        if row_i % 2 == 0:
            for c in [ws_sum.cell(row=row_i, column=2), cell_val]:
                c.fill = PatternFill("solid", fgColor=ROW_ALT)
        for c in [ws_sum.cell(row=row_i, column=2), cell_val]:
            _border_all(c)

    ws_sum.column_dimensions["B"].width = 30
    ws_sum.column_dimensions["C"].width = 20

    # CoB breakdown
    from collections import defaultdict
    cob_map: dict[str, dict[str, float]] = defaultdict(lambda: {"count": 0, "lrc": 0, "lic": 0, "icl": 0})
    for r in batch_results:
        cob = r.get("cob", "?")
        cob_map[cob]["count"] += 1
        cob_map[cob]["lrc"] += r.get("lrc_net", 0)
        cob_map[cob]["lic"] += r.get("lic", 0)
        cob_map[cob]["icl"] += r.get("insurance_contract_liability", r.get("total_liability", 0))

    start_row = 14
    ws_sum.merge_cells(f"B{start_row}:F{start_row}")
    ws_sum.cell(row=start_row, column=2, value="Breakdown Per Line of Business (COB)")
    _header_style(ws_sum.cell(row=start_row, column=2), color=HEADER_LIGHT)

    cob_headers = ["COB", "Jml Polis", "LRC Neto", "LIC", "Total ICL"]
    for ci, h in enumerate(cob_headers, 2):
        cell = ws_sum.cell(row=start_row + 1, column=ci, value=h)
        _header_style(cell, color=HEADER_LIGHT)

    for ri, (cob, vals) in enumerate(cob_map.items(), start_row + 2):
        row_data = [cob, int(vals["count"]), vals["lrc"], vals["lic"], vals["icl"]]
        for ci, v in enumerate(row_data, 2):
            c = ws_sum.cell(row=ri, column=ci, value=round(v, 2) if isinstance(v, float) else v)
            if isinstance(v, float):
                c.number_format = '#,##0.00'
            _border_all(c)
            if ri % 2 == 0:
                c.fill = PatternFill("solid", fgColor=ROW_ALT)

    # ── Sheet 2: Detail Per Polis ──────────────────────────────────
    ws_det = wb.create_sheet("Detail Per Polis")

    det_cols = [
        ("No", 5), ("Policy No", 20), ("COB", 8), ("Metode", 8),
        ("TSI (IDR)", 16), ("GWP (IDR)", 16), ("PVFCF (IDR)", 18),
        ("RA (IDR)", 14), ("CSM (IDR)", 14), ("LRC Neto (IDR)", 18),
        ("LIC (IDR)", 16), ("ICL / Total (IDR)", 20),
        ("Onerous", 10), ("Trace ID", 38),
    ]

    for col_i, (header, width) in enumerate(det_cols, 1):
        ws_det.column_dimensions[get_column_letter(col_i)].width = width
        cell = ws_det.cell(row=1, column=col_i, value=header)
        _header_style(cell)

    ws_det.row_dimensions[1].height = 25

    for row_i, r in enumerate(batch_results, 2):
        method = r.get("method", "PAA")
        is_onerous = r.get("is_onerous", False)

        row_vals = [
            row_i - 1,
            r.get("policyno", r.get("contract_id", "")),
            r.get("cob", ""),
            method,
            r.get("tsi", 0),
            r.get("gwp", 0),
            r.get("pvfcf", 0),
            r.get("risk_adjustment", r.get("ra", 0)),
            r.get("csm", 0),
            r.get("lrc_net", 0),
            r.get("lic", 0),
            r.get("insurance_contract_liability", r.get("total_liability", 0)),
            "YA" if is_onerous else "TIDAK",
            r.get("trace_id", ""),
        ]

        for col_i, val in enumerate(row_vals, 1):
            c = ws_det.cell(row=row_i, column=col_i, value=round(val, 2) if isinstance(val, float) else val)
            if isinstance(val, float) and col_i >= 5:
                c.number_format = '#,##0.00'
            _border_all(c)
            if row_i % 2 == 0:
                c.fill = PatternFill("solid", fgColor=ROW_ALT)

        # Warna merah jika onerous
        if is_onerous:
            for col_i in range(1, len(det_cols) + 1):
                ws_det.cell(row=row_i, column=col_i).font = Font(color=ACCENT_RED)

    # Freeze header
    ws_det.freeze_panes = "A2"

    # ── Sheet 3: Sistem Pembukuan GMM ────────────────────────────
    ws_gmm = wb.create_sheet("Sistem Pembukuan GMM")

    # --- Collect GMM results (up to 3 skenario) ---
    gmm_results = [r for r in batch_results if r.get("method") in ("GMM", "BBA") and r.get("gmm_detail")]
    # Use up to first 3 as Initial / First Change / Second Change
    scenarios = gmm_results[:3]
    scenario_labels = ["BS dan PL Initial", "BS dan PL First Change", "BS dan PL Second Change"]

    # --- Styling helpers ---
    bold_font = Font(bold=True, size=10)
    bold_white = Font(bold=True, size=10, color=WHITE)
    normal_font = Font(size=10)
    header_fill = PatternFill("solid", fgColor=HEADER_BLUE)
    sub_header_fill = PatternFill("solid", fgColor=HEADER_LIGHT)
    total_fill = PatternFill("solid", fgColor="D6E4F0")
    num_fmt = '#,##0.00'
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"),
        right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"),
        bottom=Side(style="thin", color="CCCCCC"),
    )

    def _w(row, col, value, font=None, fill=None, is_num=False):
        """Write a cell with styling."""
        c = ws_gmm.cell(row=row, column=col, value=value)
        if font:
            c.font = font
        if fill:
            c.fill = fill
        if is_num and isinstance(value, (int, float)):
            c.number_format = num_fmt
        c.border = thin_border
        return c

    # Each scenario occupies 6 columns: (BS: label+value) + empty + (PL: label+value) = A,B, gap, D,E
    # Template layout per skenario group = 6 cols:
    #   Col 1: BS label, Col 2: BS value,  gap(col 3),  Col 4: PL label, Col 5: PL value, gap(col 6)
    SCENARIO_WIDTH = 6

    # --- Row 1: Scenario group headers ---
    for si, label in enumerate(scenario_labels):
        base_col = 1 + si * SCENARIO_WIDTH
        ws_gmm.merge_cells(
            start_row=1, start_column=base_col,
            end_row=1, end_column=base_col + SCENARIO_WIDTH - 2
        )
        _w(1, base_col, label, font=bold_white, fill=header_fill)
        ws_gmm.cell(row=1, column=base_col).alignment = Alignment(horizontal="center")

    # --- Row 3: Sub-headers ---
    for si in range(len(scenario_labels)):
        bc = 1 + si * SCENARIO_WIDTH
        _w(3, bc, "Balance Sheet", font=bold_white, fill=sub_header_fill)
        ws_gmm.merge_cells(start_row=3, start_column=bc, end_row=3, end_column=bc + 1)
        ws_gmm.cell(row=3, column=bc).alignment = Alignment(horizontal="center")
        _w(3, bc + 3, "Profit Loss Statement", font=bold_white, fill=sub_header_fill)
        ws_gmm.merge_cells(start_row=3, start_column=bc + 3, end_row=3, end_column=bc + 4)
        ws_gmm.cell(row=3, column=bc + 3).alignment = Alignment(horizontal="center")

    # ============================================================
    # Define rows: Balance Sheet (left) and P&L (right)
    # We write ALL rows with labels first, then fill values per scenario
    # ============================================================

    # --- Balance Sheet rows (starting row 4) ---
    bs_rows = [
        # (row, label, detail_key, is_total, is_section_header)
        (4, "Asset", None, False, True),
        (6, "bank/piutang", "bank_piutang", False, False),
        (8, "EXPECTED INSURANCE CLAIM CEDED TO REINSURER - ASSET", "ri_claim_asset", False, False),
        (9, "EXPECTED ACQUISITION COST CEDED FROM REINSURER - ASSET", "ri_acq_cost_asset", False, False),
        (10, "EXPECTED EXPENSE CEDED FROM REINSURER - ASSET", "ri_expense_asset", False, False),
        (11, "NPR", "ri_npr", False, False),
        (12, "RISK ADJUSTMENT CEDED TO REINSURER - ASSET", "ri_ra_asset", False, False),
        (13, "CONTRACTUAL SERVICE MARGIN CEDED TO REINSURER - ASSET", "ri_csm_asset", False, False),
        (14, "LOSS COMPONENT CEDED TO REINSURER - ASSET", "ri_lc_asset", False, False),
        (15, "ACQUISITION COST DITERIMA DI MUKA FROM RI", "ri_acq_cost_dimuka", False, False),
        (16, "Total RI Asset", "total_ri_asset", True, False),
        (18, "Total Asset", "total_asset", True, False),
        (20, "Liabilty of Remaining Coverage", None, False, True),
        (21, "EXPECTED CLAIMS - LIABILITY", "expected_claims", False, False),
        (22, "EXPECTED ACQUISITION COST - LIABILITY", "expected_acq_cost", False, False),
        (23, "EXPECTED EXPENSE - LIABILITY", "expected_expense", False, False),
        (24, "RISK ADJUSTMENT - LIABILITY", "risk_adjustment_liability", False, False),
        (25, "CONTRACTUAL SERVICE MARGIN - LIABILITY", "csm_liability", False, False),
        (26, "LOSS COMPONENT - LIABILITY", "loss_component", False, False),
        (27, "DEFERRED ACQUISITION COST", "dac", False, False),
        (28, "Total Liabilty", "total_liability", True, False),
        (30, "Equity", "equity", True, False),
        (32, "Total Liability +Equity", None, True, False),  # computed
    ]

    # --- P&L rows (starting row 4, in PL column) ---
    pl_rows = [
        (4, "Insurance Revenue:", None, False, True),
        (5, "CHANGE IN EXPECTED INSURANCE CLAIM", "rev_expected_claim", False, False),
        (6, "CHANGE IN EXPECTED ACQUISITION COSTS", "rev_acq_cost", False, False),
        (7, "CHANGE IN EXPECTED INSURANCE EXPENSE", "rev_expected_expense", False, False),
        (8, "CHANGE IN RISK ADJUSTMENT", "rev_risk_adjustment", False, False),
        (9, "CHANGE IN CONTRACTUAL SERVICE MARGIN", "rev_csm", False, False),
        (10, "Total INSURANCE CONTRACT REVENUE", "total_insurance_revenue", True, False),
        (12, "Insurance Service Expense:", None, False, True),
        (13, "Initial Recognition of LC", "initial_recognition_lc", False, False),
        (14, "CHANGES RELATING TO FUTURE SERVICES PREMIUM (LC)", "changes_future_services_lc", False, False),
        (15, "INCURRED INSURANCE CLAIMS", "incurred_claims", False, False),
        (16, "INSURANCE EXPENSE", "insurance_expense", False, False),
        (17, "ACQUISITION COSTS", "acquisition_costs_expense", False, False),
        (18, "Total INSURANCE SERVICE EXPENSE", "total_insurance_expense", True, False),
        (19, "Insurance Service Result", "insurance_service_result", True, False),
        (21, "Insurance Finance Expense:", None, False, True),
        (22, "Akresian Expected Claim", "accretion_claims", False, False),
        (23, "Akresian Expected Expense", "accretion_expense", False, False),
        (24, "Akresian Risk Adjustment", "accretion_ra", False, False),
        (25, "Akresian CSM", "accretion_csm", False, False),
        (26, "Akresian LC", "accretion_lc", False, False),
        (27, "Akresian Acq Cost", "accretion_acq_cost", False, False),
        (28, "Total Akresian", "total_accretion", True, False),
        (29, "change in disc rate Expected Claim", "chg_disc_claim", False, False),
        (30, "change in disc rate Expected Expense", "chg_disc_expense", False, False),
        (31, "change in disc rate Risk Adjustment", "chg_disc_ra", False, False),
        (32, "change in disc rate CSM", "chg_disc_csm", False, False),
        (33, "change in disc rate LC", "chg_disc_lc", False, False),
        (34, "change in disc rate Acq cost", "chg_disc_acq_cost", False, False),
        (35, "Total change in disc rate", "total_chg_disc", True, False),
        (36, "Total Insurance Finance Expense", "total_finance_expense", True, False),
        (37, "Surplus/Defisit Gross", "surplus_gross", True, False),
        (39, "Insurance Ceded to Reinsurer Revenue:", None, False, True),
        (40, "CHANGE IN EXPECTED CLAIM CEDED TO REINSURER", "ri_rev_claim", False, False),
        (41, "EXPECTED ACQUISITION COST CEDED FROM REINSURER", "ri_rev_acq_cost", False, False),
        (42, "CHANGE IN EXPECTED EXPENSE CEDED FROM REINSURER", "ri_rev_expense", False, False),
        (43, "CHANGE IN NPR", "ri_rev_npr", False, False),
        (44, "CHANGE IN RISK ADJUSTMENT CEDED TO REINSURER", "ri_rev_ra", False, False),
        (45, "CHANGE IN CONTRACTUAL SERVICE MARGIN CEDED TO REINSURER", "ri_rev_csm", False, False),
        (46, "Total Insurance Ceded to Reinsurer Revenue", "total_ri_revenue", True, False),
        (48, "Insurance Ceded to Reinsurer Expense:", None, False, True),
        (49, "Initial Recognition of LC Ceded to Reinsurer", "ri_exp_lc_initial", False, False),
        (50, "CHANGES RELATING TO FUTURE SERVICES PREMIUM (LC) Ceded to Reinsurer", "ri_exp_lc_changes", False, False),
        (51, "INCURRED INSURANCE CEDED TO REINSURER CLAIMS", "ri_exp_incurred_claims", False, False),
        (52, "INSURANCE CEDED FROM REINSURER EXPENSES", "ri_exp_expense", False, False),
        (53, "INSURANCE CEDED FROM REINSURER ACQUISITION COST", "ri_exp_acq_cost", False, False),
        (54, "Total Insurance Ceded to Reinsurer Expense", "total_ri_expense", True, False),
        (56, "Insurance Finance Expense:", None, False, True),
        (57, "Akresian Expected Claim Ceded to Reinsurer", "ri_accretion_claim", False, False),
        (58, "Akresian Expected Expense Ceded from Reinsurer", "ri_accretion_expense", False, False),
        (59, "AKRESIAN NPR", "ri_accretion_npr", False, False),
        (60, "Akresian Risk Adjustment Ceded to Reinsurer", "ri_accretion_ra", False, False),
        (61, "Akresian CSM Ceded to Reinsurer", "ri_accretion_csm", False, False),
        (62, "Akresian LC Ceded to Reinsurer", "ri_accretion_lc", False, False),
        (63, "Akresian Acq Cost", "ri_accretion_acq_cost", False, False),
        (64, "Total Akresian Ceded to Reinsurer", "total_ri_accretion", True, False),
        (65, "change in disc rate Expected Claim Ceded to Reinsurer", "ri_chg_disc_claim", False, False),
        (66, "change in disc rate Expected Expense Ceded from Reinsurer", "ri_chg_disc_expense", False, False),
        (67, "CHANGE IN DISC RATE NPR", "ri_chg_disc_npr", False, False),
        (68, "change in disc rate Risk Adjustment Ceded to Reinsurer", "ri_chg_disc_ra", False, False),
        (69, "change in disc rate CSM Ceded to Reinsurer", "ri_chg_disc_csm", False, False),
        (70, "change in disc rate LC Ceded to Reinsurer", "ri_chg_disc_lc", False, False),
        (71, "change in disc rate Acq Cost Ceded to Reinsurer", "ri_chg_disc_acq_cost", False, False),
        (72, "Total change in disc rate Ceded to Reinsurer", "total_ri_chg_disc", True, False),
        (73, "Total Insurance Finance Expense", "total_ri_finance_expense", True, False),
        (74, "Surplus/Defisit RI", "surplus_ri", True, False),
        (76, "Profit/Loss", "profit_loss", True, False),
    ]

    # --- Write labels and values for each scenario ---
    for si, scenario_label in enumerate(scenario_labels):
        bc = 1 + si * SCENARIO_WIDTH  # base column for this scenario
        pl_label_col = bc + 3
        pl_val_col = bc + 4

        # Get detail data (or empty dict if no scenario)
        d = scenarios[si].get("gmm_detail", {}) if si < len(scenarios) else {}

        # Write BS labels + values
        for (row, label, key, is_total, is_header) in bs_rows:
            font = bold_font if (is_total or is_header) else normal_font
            fill = total_fill if is_total else None
            _w(row, bc, label, font=font, fill=fill)
            if key and d:
                val = d.get(key, 0)
                _w(row, bc + 1, round(val, 2) if isinstance(val, float) else val,
                   font=font, fill=fill, is_num=True)
            elif row == 32 and d:
                # Total Liability + Equity = computed
                tl = d.get("total_liability", 0)
                eq = d.get("equity", 0)
                _w(row, bc + 1, round(tl + eq, 2), font=font, fill=fill, is_num=True)
            elif row == 20 and d:
                # Section header "Liability" tag
                _w(row, bc + 1, "Liability", font=bold_font)

        # Write PL labels + values  
        for (row, label, key, is_total, is_header) in pl_rows:
            font = bold_font if (is_total or is_header) else normal_font
            fill = total_fill if is_total else None
            _w(row, pl_label_col, label, font=font, fill=fill)
            if key and d:
                val = d.get(key, 0)
                _w(row, pl_val_col, round(val, 2) if isinstance(val, float) else val,
                   font=font, fill=fill, is_num=True)

    # --- Column widths ---
    for si in range(3):
        bc = 1 + si * SCENARIO_WIDTH
        ws_gmm.column_dimensions[get_column_letter(bc)].width = 52      # BS label
        ws_gmm.column_dimensions[get_column_letter(bc + 1)].width = 18  # BS value
        ws_gmm.column_dimensions[get_column_letter(bc + 2)].width = 2   # gap
        ws_gmm.column_dimensions[get_column_letter(bc + 3)].width = 52  # PL label
        ws_gmm.column_dimensions[get_column_letter(bc + 4)].width = 18  # PL value
        ws_gmm.column_dimensions[get_column_letter(bc + 5)].width = 2   # gap

    ws_gmm.freeze_panes = "A2"

    # Output ke bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
