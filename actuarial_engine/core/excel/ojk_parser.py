"""
OJK Code Parser – Membaca SANDI LINI USAHA BERDASARKAN OJK APOLLO.

Sheet: COB
Kolom: Sandi (int), Lini Usaha (str)
"""
from __future__ import annotations

from typing import Any
import io
import logging

import openpyxl

from actuarial_engine.models.ojk_codes import OJKCode, OJKCodeTable

logger = logging.getLogger(__name__)


def parse_ojk_codes(file_bytes: bytes) -> tuple[OJKCodeTable, list[str]]:
    """
    Parse file SANDI LINI USAHA.

    Returns:
        Tuple of (OJKCodeTable, errors list).
    """
    table = OJKCodeTable()
    errors: list[str] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)

        # Find the COB sheet
        sheet_name = next(
            (s for s in wb.sheetnames if "cob" in s.lower()),
            wb.sheetnames[0]
        )
        ws = wb[sheet_name]

        # Find header row with "Sandi"
        sandi_col = None
        lini_col = None
        header_row = 0

        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True)):
            for j, val in enumerate(row):
                val_str = str(val).strip().lower() if val else ""
                if val_str == "sandi":
                    sandi_col = j
                    header_row = i + 1
                elif val_str == "lini usaha":
                    lini_col = j

        if sandi_col is None or lini_col is None:
            errors.append("Kolom 'Sandi' atau 'Lini Usaha' tidak ditemukan")
            wb.close()
            return table, errors

        # Read data rows
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            sandi_val = row[sandi_col] if sandi_col < len(row) else None
            lini_val = row[lini_col] if lini_col < len(row) else None

            if sandi_val is None or lini_val is None:
                continue

            try:
                sandi = int(float(str(sandi_val)))
            except (ValueError, TypeError):
                continue

            table.codes.append(OJKCode(
                sandi=sandi,
                lini_usaha=str(lini_val).strip(),
            ))

        wb.close()
        logger.info(f"Parsed {len(table.codes)} OJK codes")

    except Exception as e:
        errors.append(f"Error parsing OJK codes: {e}")
        logger.exception(f"Error parsing OJK codes: {e}")

    return table, errors


def parse_ojk_codes_from_path(file_path: str) -> tuple[OJKCodeTable, list[str]]:
    """Convenience: parse from file path."""
    with open(file_path, "rb") as f:
        return parse_ojk_codes(f.read())
