"""Baca detail setiap sheet dari template Excel untuk gap analysis."""
import openpyxl

def read_sheet_detail(path, sheet_name):
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet_name not in wb.sheetnames:
        print(f"Sheet '{sheet_name}' not found in {path}")
        wb.close()
        return
    ws = wb[sheet_name]
    print(f"\n=== [{sheet_name}] ({ws.max_row}r x {ws.max_column}c) ===")
    for r in range(1, min(20, ws.max_row+1)):
        row_vals = [ws.cell(r,c).value for c in range(1, min(30,ws.max_column+1))]
        non_null = [(i+1, str(v)[:40]) for i, v in enumerate(row_vals) if v is not None]
        if non_null:
            print(f"  R{r}: {non_null}")
    wb.close()


GMM_PATH = r"d:\Project\PSAK117\template\Template LRC GMM 18062026 test.xlsx"
PAA_PATH = r"d:\Project\PSAK117\template\Template PAA with accretion + LC GMM Check HE 14.04.26 (test).xlsx"

# Read LRC GMM sheets
print("\n##### TEMPLATE LRC GMM - All Sheets #####")
wb = openpyxl.load_workbook(GMM_PATH, data_only=True)
print("Sheets:", wb.sheetnames)
wb.close()

for sname in ["LRC Initial", "INIT_FCF", "Sistem Pembukuan GMM"]:
    read_sheet_detail(GMM_PATH, sname)

# Read PAA sheets
print("\n\n##### TEMPLATE PAA - All Sheets #####")
wb2 = openpyxl.load_workbook(PAA_PATH, data_only=True)
print("Sheets:", wb2.sheetnames)
wb2.close()

for sname in wb2.sheetnames if False else []:
    pass

wb2 = openpyxl.load_workbook(PAA_PATH, data_only=True)
for sname in wb2.sheetnames[:4]:
    ws = wb2[sname]
    print(f"\n=== [{sname}] ({ws.max_row}r x {ws.max_column}c) ===")
    for r in range(1, min(15, ws.max_row+1)):
        row_vals = [ws.cell(r,c).value for c in range(1, min(25,ws.max_column+1))]
        non_null = [(i+1, str(v)[:35]) for i, v in enumerate(row_vals) if v is not None]
        if non_null:
            print(f"  R{r}: {non_null}")
wb2.close()
