"""Analyze the structure of the template data files."""
import openpyxl
import os

TEMPLATES_DIR = r"d:\Web\Actuarial Engine\data\templates"

def analyze_asumsi():
    path = os.path.join(TEMPLATES_DIR, "ASUMSI 2020 - 2025 - Update.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    print("=" * 60)
    print("ASUMSI 2020 - 2025 - Update.xlsx")
    print("=" * 60)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\nSheet: {name} ({ws.max_row} rows x {ws.max_column} cols)")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=True)):
            vals = [str(v)[:50] if v is not None else "" for v in row]
            print(f"  R{i+1}: {' | '.join(vals)}")
    wb.close()

def analyze_sandi():
    path = os.path.join(TEMPLATES_DIR, "SANDI LINI USAHA BERDASARKAN OJK APOLLO - KIRIM.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    print("\n" + "=" * 60)
    print("SANDI LINI USAHA")
    print("=" * 60)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"\nSheet: {name} ({ws.max_row} rows x {ws.max_column} cols)")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True)):
            vals = [str(v)[:60] if v is not None else "" for v in row]
            print(f"  R{i+1}: {' | '.join(vals)}")
    wb.close()

def analyze_inforce(filename, label):
    path = os.path.join(TEMPLATES_DIR, filename)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("\n" + "=" * 60)
    print(f"INFORCE: {label}")
    print("=" * 60)
    print(f"Sheets: {wb.sheetnames}")
    for name in wb.sheetnames[:5]:
        ws = wb[name]
        print(f"\n--- Sheet: {name} ---")
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True)):
            vals = [str(v)[:35] if v is not None else "" for v in list(row)[:25]]
            print(f"  R{i+1}: {' | '.join(vals)}")
        # count rows
        rc = 0
        for _ in ws.iter_rows(min_row=1, values_only=True):
            rc += 1
            if rc > 200000:
                break
        print(f"  Total rows: {rc}")
    wb.close()

if __name__ == "__main__":
    analyze_sandi()
    analyze_asumsi()
    print("\n\nNow analyzing INFORCE 2023 (large file, may take a moment)...")
    analyze_inforce("01. INFORCE 31 DESEMBER 2023 - KIRIM.xlsx", "2023")
