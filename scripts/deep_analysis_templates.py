"""Deep analysis template PSAK117 untuk gap analysis."""
import openpyxl

GMM_PATH = r"d:\Project\PSAK117\template\Template LRC GMM 18062026 test.xlsx"
PAA_PATH = r"d:\Project\PSAK117\template\Template PAA with accretion + LC GMM Check HE 14.04.26 (test).xlsx"

wb = openpyxl.load_workbook(GMM_PATH, data_only=True)
print("=== LRC GMM Sheets:", wb.sheetnames)

# Sheet 1: LRC Initial
ws = wb["LRC Initial"]
print(f"\n[LRC Initial] {ws.max_row}r x {ws.max_column}c")
for r in range(1, min(30, ws.max_row+1)):
    row = [ws.cell(r,c).value for c in range(1, min(20,ws.max_column+1))]
    non_null = [(i+1, str(v)[:50]) for i, v in enumerate(row) if v is not None]
    if non_null:
        print(f"  R{r}: {non_null}")

# Cari sheet untuk Sistem Pembukuan GMM
for sn in wb.sheetnames:
    if "Sistem" in sn or "Pembukuan" in sn or "GMM" in sn:
        ws2 = wb[sn]
        print(f"\n[{sn}] {ws2.max_row}r x {ws2.max_column}c")
        for r in range(1, min(25, ws2.max_row+1)):
            row = [ws2.cell(r,c).value for c in range(1, min(25,ws2.max_column+1))]
            non_null = [(i+1, str(v)[:50]) for i, v in enumerate(row) if v is not None]
            if non_null:
                print(f"  R{r}: {non_null}")
wb.close()

print("\n\n=== PAA TEMPLATE ===")
wb2 = openpyxl.load_workbook(PAA_PATH, data_only=True)
print("PAA Sheets:", wb2.sheetnames)

for sn in wb2.sheetnames:
    ws = wb2[sn]
    print(f"\n[{sn}] {ws.max_row}r x {ws.max_column}c")
    for r in range(1, min(20, ws.max_row+1)):
        row = [ws.cell(r,c).value for c in range(1, min(20,ws.max_column+1))]
        non_null = [(i+1, str(v)[:50]) for i, v in enumerate(row) if v is not None]
        if non_null:
            print(f"  R{r}: {non_null}")
wb2.close()
