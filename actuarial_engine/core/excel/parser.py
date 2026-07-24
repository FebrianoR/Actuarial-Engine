"""
Excel Parser – Membaca 4 template Excel PSAK 117

Template yang didukung:
  00-Templateupload.xlsx  → data polis (datapolis sheet)
  01-templateassumption.xlsx → asumsi (assumsi sheet)
  02-templatetabphei.xlsx → yield curve (lrc_tabphei sheet)
  03-templatelapse.xlsx   → lapse table (lrc_lapsemmtable sheet)
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any
import io

import openpyxl
import pandas as pd


# ──────────────────────────────────────────────
#  Data classes for parsed results
# ──────────────────────────────────────────────

@dataclass
class PolicyRecord:
    """Satu baris dari sheet datapolis (00-Templateupload)."""
    sguser: str = ""
    id: str = ""
    batchnr: str = ""
    cob: str = ""           # Line of Business (A, B, C, ...)
    product: str = ""
    sob: str = ""
    policyno: str = ""
    begindt: date | None = None
    enddt: date | None = None
    valdate: date | None = None
    stspolis: str = "A"     # A = Aktif
    tsi: float = 0.0        # Total Sum Insured
    gwp: float = 0.0        # Gross Written Premium
    discount: float = 0.0
    outgo: float = 0.0
    ngwp: float = 0.0       # Net GWP
    smethode: str = "PAA"   # PAA atau GMM


@dataclass
class AssumptionRow:
    """Satu baris asumsi dari sheet assumsi (01-templateassumption)."""
    sguser: str = ""
    cob: str = ""
    syear: int = 2024
    expclaim: float = 0.0       # Expected Claim Ratio
    expopex: float = 0.0        # Expected Opex (% of premium)
    expriskadj: float = 0.0     # Expected Risk Adjustment (% of claim)


@dataclass
class YieldCurveRow:
    """Satu baris yield curve dari tabphei (02-templatetabphei)."""
    sguser: str = ""
    syear: int = 0
    sdate: date | None = None
    srate: float = 0.0          # Rate (dalam persen, mis 6.0784 = 6.0784%)


@dataclass
class LapseRow:
    """Satu baris lapse table dari 03-templatelapse."""
    sguser: str = ""
    syear: int = 1
    smonth: int = 0
    smonth_th: int = 1
    cob: str = ""
    srate: float = 0.0          # Lapse rate per bulan


@dataclass
class ParsedTemplates:
    """Kumpulan hasil parsing semua template."""
    policies: list[PolicyRecord] = field(default_factory=list)
    assumptions: list[AssumptionRow] = field(default_factory=list)
    yield_curve: list[YieldCurveRow] = field(default_factory=list)
    lapse_table: list[LapseRow] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


# ──────────────────────────────────────────────
#  Parser functions
# ──────────────────────────────────────────────

def _to_date(val: Any) -> date | None:
    """Convert berbagai format ke date."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
            try:
                return datetime.strptime(val.strip(), fmt).date()
            except ValueError:
                continue
    return None


def _to_float(val: Any) -> float:
    """Safe float conversion."""
    try:
        return float(val or 0)
    except (TypeError, ValueError):
        return 0.0


def parse_portfolio(file_bytes: bytes) -> tuple[list[PolicyRecord], list[str]]:
    """
    Parse 00-Templateupload.xlsx (sheet: datapolis).
    Returns (policies, errors).
    """
    policies: list[PolicyRecord] = []
    errors: list[str] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        # Cari sheet datapolis (case-insensitive)
        sheet_name = next(
            (s for s in wb.sheetnames if "datapolis" in s.lower() or "data" in s.lower()),
            wb.sheetnames[0]
        )
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            errors.append("Sheet datapolis kosong")
            wb.close()
            return policies, errors

        # Deteksi header (baris pertama yang memiliki 'policyno' atau 'id')
        header_row = 0
        headers: list[str] = []
        for i, row in enumerate(rows):
            row_lower = [str(v).lower().strip() if v else "" for v in row]
            if "policyno" in row_lower or "id" in row_lower:
                headers = row_lower
                header_row = i
                break

        if not headers:
            # Asumsikan baris pertama adalah header
            headers = [str(v).lower().strip() if v else f"col{j}" for j, v in enumerate(rows[0])]
            header_row = 0

        def get(row: tuple, col: str) -> Any:
            try:
                idx = headers.index(col)
                return row[idx] if idx < len(row) else None
            except ValueError:
                return None

        for row_num, row in enumerate(rows[header_row + 1:], start=header_row + 2):
            if not any(row):
                continue
            policyno = get(row, "policyno")
            if not policyno:
                continue

            rec = PolicyRecord(
                sguser=str(get(row, "sguser") or ""),
                id=str(get(row, "id") or ""),
                batchnr=str(get(row, "batchnr") or ""),
                cob=str(get(row, "cob") or ""),
                product=str(get(row, "product") or ""),
                sob=str(get(row, "sob") or ""),
                policyno=str(policyno),
                begindt=_to_date(get(row, "begindt")),
                enddt=_to_date(get(row, "enddt")),
                valdate=_to_date(get(row, "valdate")),
                stspolis=str(get(row, "stspolis") or "A"),
                tsi=_to_float(get(row, "tsi")),
                gwp=_to_float(get(row, "gwp")),
                discount=_to_float(get(row, "discount")),
                outgo=_to_float(get(row, "outgo")),
                ngwp=_to_float(get(row, "ngwp")),
                smethode=str(get(row, "smethode") or "PAA").upper(),
            )
            policies.append(rec)

        wb.close()
    except Exception as e:
        errors.append(f"Error parsing portfolio: {e}")

    return policies, errors


def parse_assumptions(file_bytes: bytes) -> tuple[list[AssumptionRow], list[str]]:
    """Parse 01-templateassumption.xlsx (sheet: assumsi)."""
    rows_out: list[AssumptionRow] = []
    errors: list[str] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet_name = next(
            (s for s in wb.sheetnames if "assum" in s.lower()),
            wb.sheetnames[0]
        )
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        headers = [str(v).lower().strip() if v else "" for v in rows[0]]

        def get(row: tuple, col: str) -> Any:
            try:
                idx = headers.index(col)
                return row[idx] if idx < len(row) else None
            except ValueError:
                return None

        for row in rows[1:]:
            if not any(row):
                continue
            sguser = get(row, "sguser")
            if not sguser:
                continue
            rows_out.append(AssumptionRow(
                sguser=str(sguser),
                cob=str(get(row, "cob") or ""),
                syear=int(_to_float(get(row, "syear")) or 2024),
                expclaim=_to_float(get(row, "expclaim")),
                expopex=_to_float(get(row, "expopex")),
                expriskadj=_to_float(get(row, "expriskadj")),
            ))

        wb.close()
    except Exception as e:
        errors.append(f"Error parsing assumptions: {e}")

    return rows_out, errors


def parse_yield_curve(file_bytes: bytes) -> tuple[list[YieldCurveRow], list[str]]:
    """Parse 02-templatetabphei.xlsx (sheet: lrc_tabphei)."""
    rows_out: list[YieldCurveRow] = []
    errors: list[str] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(v).lower().strip() if v else "" for v in rows[0]]

        def get(row: tuple, col: str) -> Any:
            try:
                idx = headers.index(col)
                return row[idx] if idx < len(row) else None
            except ValueError:
                return None

        for row in rows[1:]:
            if not any(row):
                continue
            rows_out.append(YieldCurveRow(
                sguser=str(get(row, "sguser") or ""),
                syear=int(_to_float(get(row, "syear"))),
                sdate=_to_date(get(row, "sdate")),
                srate=_to_float(get(row, "srate")),
            ))

        wb.close()
    except Exception as e:
        errors.append(f"Error parsing yield curve: {e}")

    return rows_out, errors


def parse_lapse(file_bytes: bytes) -> tuple[list[LapseRow], list[str]]:
    """Parse 03-templatelapse.xlsx (sheet: lrc_lapsemmtable)."""
    rows_out: list[LapseRow] = []
    errors: list[str] = []

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(v).lower().strip() if v else "" for v in rows[0]]

        def get(row: tuple, col: str) -> Any:
            try:
                idx = headers.index(col)
                return row[idx] if idx < len(row) else None
            except ValueError:
                return None

        for row in rows[1:]:
            if not any(row):
                continue
            rows_out.append(LapseRow(
                sguser=str(get(row, "sguser") or ""),
                syear=int(_to_float(get(row, "syear")) or 1),
                smonth=int(_to_float(get(row, "smonth"))),
                smonth_th=int(_to_float(get(row, "smonth_th")) or 1),
                cob=str(get(row, "cob") or ""),
                srate=_to_float(get(row, "srate")),
            ))

        wb.close()
    except Exception as e:
        errors.append(f"Error parsing lapse: {e}")

    return rows_out, errors
